"""LIFE_ETHICS 최종 이관 패키지 V4 export (분석·재판정 없음, V1/V2/V3 보존, 결정적 출력).

선정: data/phase2/LIFE_ETHICS/v4_strict_audit/v4_final_content_asset_index.jsonl
원문: Phase 1 data/questions/json/*.json, data/questions/images/*.png (읽기 전용)
출력: FINAL_TRANSFER/LIFE_ETHICS_V4/{manifest.json, README.md, questions.jsonl, questions.xlsx, images/}
      FINAL_TRANSFER/LIFE_ETHICS_FINAL_TRANSFER_V4.zip (루트 LIFE_ETHICS_V4/)
결정성: xlsx 문서 속성·zip 엔트리 시각을 고정해 재실행 시 byte-identical 출력.
"""
from __future__ import annotations

import collections
import datetime as dt
import hashlib
import io
import json
import re
import shutil
import sqlite3
import zipfile
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
P2 = ROOT / "data/phase2/LIFE_ETHICS"
AUD = P2 / "v4_strict_audit"
INDEX = AUD / "v4_final_content_asset_index.jsonl"
BASE = ROOT / "FINAL_TRANSFER"
OUT = BASE / "LIFE_ETHICS_V4"
ZIP = BASE / "LIFE_ETHICS_FINAL_TRANSFER_V4.zip"
EXAM_PERIOD = {"JUNE": "6월 모의평가", "SEPTEMBER": "9월 모의평가", "CSAT": "대학수학능력시험"}
FIXED_DT = (2026, 1, 1, 0, 0, 0)


def sha256(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def normalized_zip(entries: list[tuple[str, bytes]]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries:
            zi = zipfile.ZipInfo(name, date_time=FIXED_DT)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            z.writestr(zi, data)
    return buf.getvalue()


def main():
    idx = [json.loads(l) for l in open(INDEX, encoding="utf-8")]
    ids_idx = [r["question_id"] for r in idx]
    comp = collections.Counter(r["source"] for r in idx)
    aud = [json.loads(l) for l in open(AUD / "v4_strict_audit.jsonl", encoding="utf-8")]
    corr_drop = {r["question_id"] for r in aud if r["final_decision"] == "DROP"}
    n = len(aud) - len(corr_drop)
    if not (len(idx) == len(set(ids_idx)) == n):
        raise SystemExit(f"STOP: selection count mismatch {len(idx)} vs expected {n}")
    all_ids = {p.stem for p in (ROOT / "data/questions/json").glob("*.json")}

    if OUT.exists():
        shutil.rmtree(OUT)
    (OUT / "images").mkdir(parents=True)

    rows, src_hash = [], {}
    for sel in sorted(idx, key=lambda r: r["question_id"]):
        qid = sel["question_id"]
        q = json.load(open(ROOT / f"data/questions/json/{qid}.json", encoding="utf-8"))
        s, o, a = q["source"], q["original"], q["answer"]
        data = (ROOT / o["question_image"]).read_bytes()
        src_hash[qid] = sha256(data)
        (OUT / "images" / f"{qid}.png").write_bytes(data)
        rows.append({
            "question_id": qid,
            "subject": s["subject"],
            "exam_year": s["exam_year"],
            "academic_year": s["academic_year"],
            "exam_session": s["exam_type"],
            "exam_name": EXAM_PERIOD.get(s["exam_type"]),
            "exam_date": None,
            "question_number": s["question_number"],
            "stem": o["stem"],
            "passage": o["passage"],
            "choices": o["choices"],
            "question_text": o["raw_text"],
            "official_answer": a["value"],
            "points": o["points"],
            "source_pdf": s["split_pdf"],
            "source_page": s["split_pdf_page"],
            "source_compilation": s["original_compilation"],
            "original_page": s["original_compilation_page"],
            "image": f"images/{qid}.png",
            "image_sha256": src_hash[qid],
            "final_asset_source": sel["source"],
            "review_required": sel["review_required"],
            "primary_curriculum_code": sel["primary_curriculum_code"],
            "kice_sample_link": sel["kice_sample_link"],
            "v4_basis": sel["v4_basis"],
            "adaptation_required": sel["adaptation_required"],
        })

    with open(OUT / "questions.jsonl", "w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")

    wb = Workbook()
    wb.properties.creator = "LIFE_ETHICS_V4"
    wb.properties.created = wb.properties.modified = dt.datetime(*FIXED_DT)
    ws = wb.active
    ws.title = "LIFE_ETHICS_V4"
    head = ["question_id", "exam_year", "exam_session", "question_number", "official_answer", "points",
            "final_asset_source", "review_required", "primary_curriculum_code", "kice_sample_link",
            "v4_basis", "adaptation_required", "image_filename"]
    ws.append(head)
    for r in rows:
        link = r["kice_sample_link"]
        ws.append([r["question_id"], r["exam_year"], r["exam_session"], r["question_number"], r["official_answer"],
                   r["points"], r["final_asset_source"], r["review_required"], r["primary_curriculum_code"],
                   link["sample_question"] if link else None,
                   r["v4_basis"], r["adaptation_required"], r["image"]])
    for i, w in enumerate([32, 10, 12, 10, 10, 7, 28, 10, 18, 22, 16, 12, 44], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    raw = io.BytesIO()
    wb.save(raw)
    fixed_w3c = dt.datetime(*FIXED_DT).strftime("%Y-%m-%dT%H:%M:%SZ")
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as zx:  # xlsx 내부 zip 시각·저장 시각(modified) 정규화
        parts = []
        for i in zx.infolist():
            data = zx.read(i.filename)
            if i.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                              lambda m: m.group(1) + fixed_w3c.encode() + m.group(2), data)
            parts.append((i.filename, data))
        (OUT / "questions.xlsx").write_bytes(normalized_zip(parts))

    jl = [json.loads(l) for l in open(OUT / "questions.jsonl", encoding="utf-8")]
    pngs = sorted((OUT / "images").glob("*.png"))
    summ = json.load(open(AUD / "v4_strict_audit_summary.json", encoding="utf-8"))
    manifest = {
        "package": "LIFE_ETHICS_FINAL_TRANSFER_V4",
        "subject": "LIFE_ETHICS",
        "selection_source": "data/phase2/LIFE_ETHICS/v4_strict_audit/v4_final_content_asset_index.jsonl",
        "total_questions": len(jl),
        "composition": dict(sorted(collections.Counter(r["final_asset_source"] for r in jl).items())),
        "review_required_count": sum(r["review_required"] for r in jl),
        "official_answer_null_count": sum(r["official_answer"] is None for r in jl),
        "image_count": len(pngs),
        "selection_history": {
            "phase2c_keep": 261,
            "official_curriculum_core": 106,
            "kice_2028_restored": 30,
            "final_v2": 136,
            "final_v3": summ["base"]["v3_input"],
            "v4_reviewed": summ["reviewed"],
            "v4_drop_from_v3": summ["v3_drop"],
            "v4_civil_disobedience_policy_restore": summ["civil_disobedience_policy_restore"]["count"],
            "final_v4": len(jl),
        },
        "previous_package": "FINAL_TRANSFER/LIFE_ETHICS_V3 (historical snapshot, freeze 3a22b7f)",
        "files": {"questions": "questions.jsonl", "spreadsheet": "questions.xlsx", "images": "images/", "readme": "README.md"},
        "notes": "stem/passage/choices/question_text는 Phase 1 OCR(보조). 원문 Source of Truth는 images/<question_id>.png. "
                 "official_answer는 Phase 1 값 그대로(LIFE_ETHICS Phase 1에 공식 정답 출처 없음, 추론 없이 null 유지).",
        "status": "FROZEN_V4",
    }
    (OUT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "README.md").write_text(
        "# LIFE_ETHICS FINAL TRANSFER V4\n\n"
        f"- 문항 수: {len(jl)} (CURRICULUM_EXPLICIT_CORE {manifest['composition'].get('CURRICULUM_EXPLICIT_CORE', 0)}, "
        f"KICE_2028_RESTORED {manifest['composition'].get('KICE_2028_RESTORED', 0)}, "
        f"CIVIL_DISOBEDIENCE_POLICY_RESTORED {manifest['composition'].get('CIVIL_DISOBEDIENCE_POLICY_RESTORED', 0)})\n"
        f"- 선정 기준: 정답 결정 논리가 공식 통합사회([별책7]) 또는 KICE 2028 예시에서 실제로 확인되는 문항만 유지"
        f"(V3 {summ['base']['v3_input']}문항 전수 검토, {summ['v3_drop']}문항 제외). 시민 불복종은 정책상 "
        f"사상가 비교형도 포함({summ['civil_disobedience_policy_restore']['count']}문항 복원).\n"
        "- `official_evidence`는 data/phase2/LIFE_ETHICS/v4_strict_audit/v4_strict_audit.jsonl 참조.\n"
        "- `questions.jsonl`: 1행 = 1문항. 텍스트는 Phase 1 OCR 보조 자료이며 원문은 `images/<question_id>.png`.\n"
        "- `questions.xlsx`: 사람 확인용 목록(원문 없음).\n"
        "- `official_answer`: 공식 정답 출처가 없어 모두 null(추론하지 않음).\n"
        "- `review_required`: 후속 제작 시 사람이 주의해서 볼 문항 표시(제외 사유 아님).\n"
        "- `adaptation_required`: NONE 또는 LIGHT(통합사회 문항화 시 표현·오답 선지 경미 수정).\n",
        encoding="utf-8")

    # ---- QA ----
    xl_rows = openpyxl.load_workbook(OUT / "questions.xlsx").active.max_row - 1
    ids_jl = [r["question_id"] for r in jl]
    ids_png = [p.stem for p in pngs]
    db_ids = {r[0] for r in sqlite3.connect(ROOT / "data/database/questions.db").execute("select id from questions")}
    qa = {
        "selection_count": len(idx),
        "jsonl_count": len(jl),
        "excel_rows": xl_rows,
        "png_count": len(pngs),
        "duplicate_ids": len(ids_jl) - len(set(ids_jl)),
        "missing_ids": len(set(ids_idx) - set(ids_jl)) + len(set(ids_idx) - set(ids_png)),
        "extra_ids": len(set(ids_jl) - set(ids_idx)) + len(set(ids_png) - set(ids_idx)),
        "correction_drop_contamination": len((set(ids_jl) | set(ids_png)) & corr_drop),
        "not_in_phase1": len(set(ids_jl) - all_ids) + sum(i not in db_ids for i in ids_jl),
        "source_to_export_hash_mismatch": sum(sha256((OUT / "images" / f"{i}.png").read_bytes()) != src_hash[i] for i in ids_idx),
    }
    if not (qa["selection_count"] == qa["jsonl_count"] == qa["excel_rows"] == qa["png_count"] == n
            and all(qa[k] == 0 for k in ("duplicate_ids", "missing_ids", "extra_ids", "correction_drop_contamination",
                                         "not_in_phase1", "source_to_export_hash_mismatch"))):
        raise SystemExit(f"STOP: QA failed {qa}")

    entries = [(str(p.relative_to(BASE)), p.read_bytes()) for p in sorted(OUT.rglob("*")) if p.is_file()]
    ZIP.write_bytes(normalized_zip(entries))

    with zipfile.ZipFile(ZIP) as z:
        bad = z.testzip()
        names = z.namelist()
        zimgs = [x for x in names if x.startswith("LIFE_ETHICS_V4/images/") and x.endswith(".png")]
        zqa = {
            "corrupt_files": 0 if bad is None else 1,
            "root_dirs": sorted({x.split("/")[0] for x in names}),
            "required_files": all(f"LIFE_ETHICS_V4/{f}" in names for f in ("questions.jsonl", "questions.xlsx", "manifest.json", "README.md")),
            "images_count": len(zimgs),
            "jsonl_records": sum(1 for l in z.read("LIFE_ETHICS_V4/questions.jsonl").decode("utf-8").splitlines() if l.strip()),
            "export_to_zip_hash_mismatch": sum(sha256(z.read(x)) != src_hash[Path(x).stem] for x in zimgs),
            "entries": len(names),
        }
    if not (zqa["corrupt_files"] == 0 and zqa["root_dirs"] == ["LIFE_ETHICS_V4"] and zqa["required_files"]
            and zqa["images_count"] == zqa["jsonl_records"] == n and zqa["export_to_zip_hash_mismatch"] == 0):
        raise SystemExit(f"STOP: ZIP verification failed {zqa}")

    print(json.dumps({"manifest": {k: manifest[k] for k in ("total_questions", "composition", "review_required_count",
                                                             "official_answer_null_count", "image_count", "selection_history")},
                      "qa": qa, "zip_qa": zqa, "zip_bytes": ZIP.stat().st_size, "zip_sha256": sha256(ZIP.read_bytes())},
                     ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
