"""LIFE_ETHICS 최종 이관 패키지 export (분석·재판정 없음).

선정: data/phase2/LIFE_ETHICS/final_curation.jsonl 의 final_decision == KEEP
원문: Phase 1 data/questions/json/*.json, data/questions/images/*.png (읽기 전용)
출력: FINAL_TRANSFER/LIFE_ETHICS/{questions.jsonl, LIFE_ETHICS_FINAL_TRANSFER.xlsx, manifest.json, images/}
      FINAL_TRANSFER/LIFE_ETHICS_FINAL_TRANSFER.zip
"""
from __future__ import annotations

import collections
import hashlib
import json
import shutil
import sqlite3
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
CUR = ROOT / "data/phase2/LIFE_ETHICS/final_curation.jsonl"
BASE = ROOT / "FINAL_TRANSFER"
OUT = BASE / "LIFE_ETHICS"
ZIP = BASE / "LIFE_ETHICS_FINAL_TRANSFER.zip"
XLSX = "LIFE_ETHICS_FINAL_TRANSFER.xlsx"
EXAM_PERIOD = {"JUNE": "6월 모의평가", "SEPTEMBER": "9월 모의평가", "CSAT": "대학수학능력시험"}


def sha256(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def main():
    all_ids = sorted(p.stem for p in (ROOT / "data/questions/json").glob("*.json"))
    cur = [json.loads(l) for l in open(CUR, encoding="utf-8")]
    decided = {r["question_id"]: r["final_decision"] for r in cur}
    keep = sorted(q for q, d in decided.items() if d == "KEEP")
    counts = collections.Counter(decided.values())
    excluded = {"MECHANISM_ONLY": counts["MECHANISM_ONLY"], "DROP": counts["DROP"],
                "NOT_CANDIDATE": len(set(all_ids) - set(decided))}

    if OUT.exists():
        shutil.rmtree(OUT)
    (OUT / "images").mkdir(parents=True)

    rows, mismatches = [], []
    for qid in keep:
        q = json.load(open(ROOT / f"data/questions/json/{qid}.json", encoding="utf-8"))
        s, o, a = q["source"], q["original"], q["answer"]
        src_png = ROOT / o["question_image"]
        dst_png = OUT / "images" / f"{qid}.png"
        h0 = sha256(src_png)
        shutil.copy2(src_png, dst_png)
        if sha256(dst_png) != h0:
            mismatches.append(qid)
        rows.append({
            "question_id": qid,
            "subject": s["subject"],
            "year": s["academic_year"],
            "exam_year": s["exam_year"],
            "exam_type": s["exam_type"],
            "exam_period": EXAM_PERIOD.get(s["exam_type"]),
            "exam_date": None,
            "exam_key": s["exam_key"],
            "question_number": s["question_number"],
            "question_text": o["raw_text"],
            "stem": o["stem"],
            "passage": o["passage"],
            "choices": o["choices"],
            "official_answer": a["value"],
            "official_answer_source": a["source"],
            "official_answer_verified": a["verified"],
            "points": o["points"],
            "has_visual_material": o["has_visual_material"],
            "source_compilation": s["original_compilation"],
            "source_file": s["split_pdf"],
            "source_page": s["split_pdf_page"],
            "original_page": s["original_compilation_page"],
            "question_image": f"images/{qid}.png",
            "image_sha256": h0,
            "text_note": "question_text/choices는 OCR 결과(원문 대조용 보조). Source of Truth는 question_image.",
            "phase1_review_required": q["extraction"]["review_required"],
            "selection_source": "LIFE_ETHICS_PHASE2C",
            "selection_decision": "KEEP",
        })

    with open(OUT / "questions.jsonl", "w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")

    wb = Workbook()
    ws = wb.active
    ws.title = "LIFE_ETHICS_KEEP"
    head = ["question_id", "year", "exam_type", "exam_period", "question_number", "question_text", "choices",
            "official_answer", "has_visual_material", "source_file", "source_page", "original_page", "image_path",
            "phase1_review_required"]
    ws.append(head)
    for r in rows:
        ws.append([r["question_id"], r["year"], r["exam_type"], r["exam_period"], r["question_number"],
                   r["question_text"], "\n".join(f"{c['number']}. {c['text']}" for c in r["choices"]),
                   r["official_answer"], r["has_visual_material"], r["source_file"], r["source_page"],
                   r["original_page"], r["question_image"], r["phase1_review_required"]])
    for i, w in enumerate([32, 7, 11, 16, 9, 80, 40, 10, 10, 26, 9, 9, 44, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
    for row in ws.iter_rows(min_row=2):
        for i in (5, 6):
            row[i].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(OUT / XLSX)

    # ---- manifest (counts read back from written files) ----
    jl = [json.loads(l) for l in open(OUT / "questions.jsonl", encoding="utf-8")]
    pngs = sorted((OUT / "images").glob("*.png"))
    manifest = {
        "subject": "LIFE_ETHICS",
        "source_questions": len(all_ids),
        "selected_questions": len(jl),
        "selection_rule": "Phase2C final_decision == KEEP",
        "excluded": excluded,
        "files": {"questions": "questions.jsonl", "spreadsheet": XLSX, "images": "images/"},
        "image_count": len(pngs),
        "official_answers_linked": sum(1 for r in jl if r["official_answer"] is not None),
        "notes": "question_text/choices는 Phase 1 OCR. 원문 확인은 images/<question_id>.png. "
                 "공식 정답표 미연결 문항의 official_answer는 null(추론 없음).",
        "status": "FINAL",
    }
    json.dump(manifest, open(OUT / "manifest.json", "w", encoding="utf-8"), ensure_ascii=False, indent=2)

    # ---- QA ----
    import openpyxl
    xl_rows = openpyxl.load_workbook(OUT / XLSX).active.max_row - 1
    ids = [r["question_id"] for r in jl]
    db_ids = {r[0] for r in sqlite3.connect(ROOT / "data/database/questions.db").execute("select id from questions")}
    qa = {
        "A_counts_equal": len(keep) == len(jl) == xl_rows == len(pngs) == manifest["selected_questions"],
        "B_duplicate_ids": len(ids) - len(set(ids)),
        "C_missing_images": [r["question_image"] for r in jl if not (OUT / r["question_image"]).exists()],
        "D_not_in_phase1_db": [i for i in ids if i not in db_ids],
        "E_non_keep_exported": [i for i in ids if decided.get(i) != "KEEP"],
        "G_image_hash_mismatches": mismatches
        + [r["question_id"] for r in jl if sha256(OUT / r["question_image"]) != r["image_sha256"]],
        "self_contained_paths": all(not r["question_image"].startswith(("/", "..", "data/")) for r in jl),
        "sum_800": len(jl) + sum(excluded.values()) == len(all_ids),
    }
    assert qa["A_counts_equal"] and not qa["B_duplicate_ids"] and not qa["C_missing_images"]
    assert not qa["D_not_in_phase1_db"] and not qa["E_non_keep_exported"] and not qa["G_image_hash_mismatches"]
    assert qa["self_contained_paths"] and qa["sum_800"]

    if ZIP.exists():
        ZIP.unlink()
    with zipfile.ZipFile(ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        for p in sorted(OUT.rglob("*")):
            if p.is_file():
                z.write(p, p.relative_to(ROOT))
    with zipfile.ZipFile(ZIP) as z:
        names = z.namelist()
        assert all(n.startswith("FINAL_TRANSFER/LIFE_ETHICS/") for n in names)
        assert sum(n.endswith(".png") for n in names) == len(pngs)

    print(json.dumps({"manifest": manifest, "qa": qa, "excel_rows": xl_rows,
                      "zip": str(ZIP.relative_to(ROOT)), "zip_bytes": ZIP.stat().st_size,
                      "zip_entries": len(names)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
