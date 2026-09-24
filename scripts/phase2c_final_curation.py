"""Phase 2C — LIFE_ETHICS 최종 선별 결과 조립 (1회성).

판정 자체는 문항 OCR 원문·2A/2B 결과·2028 seed·curriculum seed를 20~30문항 단위로 읽고 사람이 내린 의미 판정이며,
data/phase2/LIFE_ETHICS/final_curation_decisions.txt 에 기록돼 있다. 이 스크립트는 그 기록을 기존 결과와 합쳐
출력 파일만 만든다(키워드 판정 없음). 기존 Phase 1/2A/2B 파일은 읽기만 한다.

출력: data/phase2/LIFE_ETHICS/final_curation.jsonl, final_curation_summary.json, LIFE_ETHICS_FINAL.xlsx
"""
from __future__ import annotations

import collections
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
P2 = ROOT / "data/phase2/LIFE_ETHICS"
DEC = {"K": "KEEP", "M": "MECHANISM_ONLY", "D": "DROP"}
EXAM_KO = {"JUNE": "6월 모의평가", "SEPTEMBER": "9월 모의평가", "CSAT": "대학수학능력시험"}


def load_jsonl(p):
    return {json.loads(l)["question_id"]: json.loads(l) for l in open(p, encoding="utf-8")}


def main():
    a2 = load_jsonl(P2 / "analysis.jsonl")
    b2 = load_jsonl(P2 / "curriculum_analysis.jsonl")
    seed = json.load(open(P2 / "curriculum_seed_2022.json", encoding="utf-8"))
    area_name = {x["area_code"]: x["area_name"] for x in seed["areas"]}
    std_codes = {s["code"] for x in seed["areas"] for s in x["standards"]}
    samples_known = {s["sample_id"] for s in json.load(open(P2 / "seed_2028.json", encoding="utf-8"))["samples"]}

    cand_a = {k for k, v in a2.items() if v["relevance"] in ("DIRECT", "ADAPTABLE")}
    cand_b = {k for k, v in b2.items() if v["curriculum_relevance"] in ("DIRECT", "ADAPTABLE")}
    cands = cand_a | cand_b

    decisions = {}
    for line in open(P2 / "final_curation_decisions.txt", encoding="utf-8"):
        if line.startswith("#") or not line.strip():
            continue
        f = line.rstrip("\n").split("|")
        assert len(f) == 8, line[:80]
        short, d, smp, stds, note, reuse, reason, rr = f
        y, e, q = short.split("_")
        qid = f"{y}_{e}_LIFE_ETHICS_{q}"
        assert qid not in decisions, qid
        decisions[qid] = (DEC[d], smp, stds, note, reuse, reason.strip(), rr == "1")
    assert set(decisions) == cands, (len(decisions), len(cands), sorted(set(decisions) ^ cands)[:5])

    recs = []
    for qid in sorted(cands):
        dec, smp, stds, note, reuse, reason, rr = decisions[qid]
        stds = [f"10통사{s.strip()}" for s in stds.split(",") if s.strip()]
        assert all(s in std_codes for s in stds), (qid, stds)
        smps = [s for s in smp.split(",") if s]
        assert all(s in samples_known for s in smps), (qid, smps)
        areas = sorted({s.rsplit("-", 1)[0] for s in stds})
        if dec == "MECHANISM_ONLY":
            assert reuse, qid
        fam = a2[qid]["mechanism_family"] if dec != "DROP" and a2[qid]["mechanism_family"] else ""
        recs.append({
            "question_id": qid,
            "final_decision": dec,
            "phase2a_relevance": a2[qid]["relevance"],
            "curriculum_relevance": b2[qid]["curriculum_relevance"] if qid in b2 else "NOT_ASSESSED(2A DIRECT/ADAPTABLE)",
            "sample_matches": smps if dec != "DROP" else [],
            "curriculum_area": [f"{c} {area_name[c]}" for c in areas] if dec != "DROP" else [],
            "achievement_standard": stds if dec != "DROP" else [],
            "mechanism_family": fam,
            "mechanism_note": note,
            "reusable_element": reuse,
            "reason": reason,
            "review_required": rr,
            "candidate_source": [s for s, c in (("2A", cand_a), ("2B", cand_b)) if qid in c],
        })

    with open(P2 / "final_curation.jsonl", "w", encoding="utf-8") as f:
        for r in recs:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")

    # ---------------- summary ----------------
    C = collections.Counter
    dec_c = C(r["final_decision"] for r in recs)

    def dist(key_fn):
        out = collections.defaultdict(C)
        for r in recs:
            k = key_fn(r)
            if k:
                out[k][r["final_decision"]] += 1
        return {k: dict(v) for k, v in sorted(out.items())}

    def per_area(dec):
        c = C(a for r in recs if r["final_decision"] == dec for a in r["curriculum_area"])
        return dict(sorted(c.items()))

    def per_fam(dec):
        c = C(r["mechanism_family"] or "(Phase 2A family 없음)" for r in recs if r["final_decision"] == dec)
        return dict(c.most_common())

    drops = [r for r in recs if r["final_decision"] == "DROP"]
    kw = C()
    for r in drops:
        for m in re.findall(r"'([^']+)'", r["reason"]):
            kw[m] += 1
    cause = C()
    rules = [("생윤 직업 윤리 사상(베버·마르크스·칼뱅·순자 등)", r"직업|베버|마르크스|칼뱅"),
             ("요나스 책임 윤리 세부", r"요나스"),
             ("동양 사상·생사관 세부", r"동양|유교|도가|불교|공자|맹자|장자|노자|묵자|순자|생사관"),
             ("관혼상제·효 등 가정 윤리 의례", r"관혼상제|의례|관례|제례|효"),
             ("윤리학 분류(메타·실천 윤리)", r"윤리학"),
             ("사회 계약론(홉스·로크) 세부", r"홉스|로크"),
             ("정의 전쟁론 단일 사상가", r"정의 전쟁론"),
             ("아도르노 문화 산업·대중문화", r"아도르노"),
             ("생명 윤리·사랑·성·예술·스포츠 등", r"생명 윤리|사랑|성과|예술|플라톤|스포츠|유전자|동물 실험"),
             ("과학자 책임·과학 기술 정책", r"과학"),
             ("해외 원조·니부어 등 기타 사상가 단일 문항", r"")]
    for r in drops:
        for name, pat in rules:
            if re.search(pat, r["reason"]):
                cause[name] += 1
                break
    keyword_fp = sum(1 for r in drops if re.search(r"잘못 연결|과다 연결|과다 포함", r["reason"]))

    non_cand = 800 - len(recs)
    total_check = dec_c["KEEP"] + dec_c["MECHANISM_ONLY"] + dec_c["DROP"] + non_cand
    summary = {
        "total_candidates": len(recs),
        "candidate_union": {"phase2a_direct_adaptable": len(cand_a), "phase2b_direct_adaptable": len(cand_b),
                            "overlap": len(cand_a & cand_b)},
        "KEEP": dec_c["KEEP"],
        "MECHANISM_ONLY": dec_c["MECHANISM_ONLY"],
        "DROP": dec_c["DROP"],
        "review_required": sum(r["review_required"] for r in recs),
        "review_required_by_decision": dict(C(r["final_decision"] for r in recs if r["review_required"])),
        "by_phase2a_source": dist(lambda r: r["phase2a_relevance"] if "2A" in r["candidate_source"] else None),
        "by_phase2b_source": dist(lambda r: r["curriculum_relevance"] if "2B" in r["candidate_source"] else None),
        "keep_by_curriculum_area": per_area("KEEP"),
        "mechanism_only_by_curriculum_area": per_area("MECHANISM_ONLY"),
        "keep_without_curriculum_area": sum(1 for r in recs if r["final_decision"] == "KEEP" and not r["curriculum_area"]),
        "keep_by_sample": dict(sorted(C(s for r in recs if r["final_decision"] == "KEEP" for s in r["sample_matches"]).items())),
        "keep_by_mechanism_family": per_fam("KEEP"),
        "mechanism_only_by_mechanism_family": per_fam("MECHANISM_ONLY"),
        "false_positive_drop_count": dec_c["DROP"],
        "false_positive_drop_keyword_misfire": keyword_fp,
        "false_positive_drop_out_of_scope_detail": dec_c["DROP"] - keyword_fp,
        "false_positive_causes": dict(cause.most_common()),
        "false_positive_trigger_keywords": dict(kw.most_common(12)),
        "dedup_policy": "유사·반복 기출은 제거하지 않음. 동일 family·동일 사상가 문항도 각각 독립 판정. 파일 중복은 Phase 1 기준 0건.",
        "life_ethics_800_check": {
            "KEEP": dec_c["KEEP"], "MECHANISM_ONLY": dec_c["MECHANISM_ONLY"], "DROP": dec_c["DROP"],
            "NOT_CANDIDATE": non_cand, "sum": total_check, "ok": total_check == 800,
        },
        "method": "후보 461문항을 20~30문항 batch로 OCR 원문·2A/2B 결과·seed와 대조해 의미 판정(pilot 30 → batch 1회). "
                  "문항별 AI 호출·신규 OCR·Phase 1/2A/2B 재실행 없음.",
    }
    assert total_check == 800
    json.dump(summary, open(P2 / "final_curation_summary.json", "w", encoding="utf-8"), ensure_ascii=False, indent=2)

    # ---------------- xlsx (800 rows) ----------------
    by_id = {r["question_id"]: r for r in recs}
    wb = Workbook()
    ws = wb.active
    ws.title = "LIFE_ETHICS_FINAL"
    head = ["question_id", "연도", "시험", "문항번호", "Phase2A 판정", "Phase2B 판정", "최종 판정", "2028 예시 대응 문항",
            "통합사회 영역", "성취기준", "mechanism family", "재사용 가능 요소", "최종 판정 이유", "review_required",
            "원문 JSON 경로", "원문 PNG 경로"]
    ws.append(head)
    fills = {"KEEP": "C6EFCE", "MECHANISM_ONLY": "FFEB9C", "DROP": "FFC7CE", "NOT_CANDIDATE": "EDEDED"}
    counts = C()
    for qid in sorted(a2):
        q = json.load(open(ROOT / f"data/questions/json/{qid}.json", encoding="utf-8"))
        src = q["source"]
        jp, pp = f"data/questions/json/{qid}.json", f"data/questions/images/{qid}.png"
        assert (ROOT / pp).exists(), pp
        r = by_id.get(qid)
        b = b2.get(qid)
        dec = r["final_decision"] if r else "NOT_CANDIDATE"
        counts[dec] += 1
        row = [qid, src["academic_year"], EXAM_KO.get(src["exam_type"], src["exam_type"]), src["question_number"],
               a2[qid]["relevance"], b["curriculum_relevance"] if b else "-", dec,
               ", ".join(r["sample_matches"]) if r else "", "; ".join(r["curriculum_area"]) if r else "",
               ", ".join(r["achievement_standard"]) if r else "", r["mechanism_family"] if r else "",
               (r["reusable_element"] if r and dec != "DROP" else ""),
               r["reason"] if r else "Phase 2A·2B 모두 DIRECT/ADAPTABLE 후보가 아니어서 Phase 2C 비대상",
               (r["review_required"] if r else False), jp, pp]
        ws.append(row)
        ws.cell(ws.max_row, 7).fill = PatternFill("solid", fgColor=fills[dec])
    assert sum(counts.values()) == 800 and counts["NOT_CANDIDATE"] == non_cand
    widths = [32, 7, 16, 8, 12, 14, 17, 12, 30, 26, 16, 50, 70, 10, 44, 44]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for i in (12, 13):
            row[i - 1].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    s2 = wb.create_sheet("Summary")
    s2.append(["항목", "값"])
    for k in ["KEEP", "MECHANISM_ONLY", "DROP"]:
        s2.append([k, counts[k]])
    s2.append(["NOT_CANDIDATE", counts["NOT_CANDIDATE"]])
    s2.append(["합계", sum(counts.values())])
    s2.append(["review_required", summary["review_required"]])
    s2.append([])
    s2.append(["KEEP 통합사회 영역", "문항 수"])
    for k, v in summary["keep_by_curriculum_area"].items():
        s2.append([k, v])
    s2.append([])
    s2.append(["DROP 원인", "문항 수"])
    for k, v in summary["false_positive_causes"].items():
        s2.append([k, v])
    s2.column_dimensions["A"].width = 46
    s2.column_dimensions["B"].width = 12
    for c in s2[1]:
        c.font = Font(bold=True)
    wb.save(P2 / "LIFE_ETHICS_FINAL.xlsx")

    print(json.dumps({k: summary[k] for k in ["total_candidates", "KEEP", "MECHANISM_ONLY", "DROP", "review_required",
                                              "life_ethics_800_check"]}, ensure_ascii=False))
    print("xlsx rows:", dict(counts))


if __name__ == "__main__":
    main()
