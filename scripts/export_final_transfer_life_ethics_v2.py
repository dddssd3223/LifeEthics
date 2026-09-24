"""LIFE_ETHICS 최종 이관 패키지 V2 export (분석·재판정 없음, V1 보존).

선정: data/phase2/LIFE_ETHICS/final_content_asset_index.jsonl (CORE 106 + KICE_2028_RESTORED 30 = 136)
원문: Phase 1 data/questions/json/*.json, data/questions/images/*.png (읽기 전용)
출력: FINAL_TRANSFER/LIFE_ETHICS_V2/{manifest.json, questions.jsonl, questions.xlsx, images/}
      FINAL_TRANSFER/LIFE_ETHICS_FINAL_TRANSFER_V2.zip (루트 LIFE_ETHICS_V2/)
"""
from __future__ import annotations

import collections
import hashlib
import json
import shutil
import sqlite3
import zipfile
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
P2 = ROOT / "data/phase2/LIFE_ETHICS"
INDEX = P2 / "final_content_asset_index.jsonl"
BASE = ROOT / "FINAL_TRANSFER"
OUT = BASE / "LIFE_ETHICS_V2"
ZIP = BASE / "LIFE_ETHICS_FINAL_TRANSFER_V2.zip"
EXAM_PERIOD = {"JUNE": "6월 모의평가", "SEPTEMBER": "9월 모의평가", "CSAT": "대학수학능력시험"}
EXPECT = {"total": 136, "CURRICULUM_EXPLICIT_CORE": 106, "KICE_2028_RESTORED": 30}


def sha256(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def main():
    idx = [json.loads(l) for l in open(INDEX, encoding="utf-8")]
    ids_idx = [r["question_id"] for r in idx]
    comp = collections.Counter(r["source"] for r in idx)
    if not (len(idx) == len(set(ids_idx)) == EXPECT["total"]
            and comp["CURRICULUM_EXPLICIT_CORE"] == EXPECT["CURRICULUM_EXPLICIT_CORE"]
            and comp["KICE_2028_RESTORED"] == EXPECT["KICE_2028_RESTORED"]):
        raise SystemExit(f"STOP: index count mismatch {len(idx)} unique={len(set(ids_idx))} {dict(comp)}")

    # 제외 대상 집합(오염 검사용)
    cur = {json.loads(l)["question_id"]: json.loads(l)["final_decision"] for l in open(P2 / "final_curation.jsonl", encoding="utf-8")}
    rc = {json.loads(l)["question_id"]: json.loads(l)["reconsideration_decision"]
          for l in open(P2 / "kice_2028_reconsideration.jsonl", encoding="utf-8")}
    all_ids = {p.stem for p in (ROOT / "data/questions/json").glob("*.json")}
    excluded = all_ids - set(ids_idx)

    if OUT.exists():
        shutil.rmtree(OUT)
    (OUT / "images").mkdir(parents=True)

    rows, src_hash = [], {}
    for sel in sorted(idx, key=lambda r: r["question_id"]):
        qid = sel["question_id"]
        q = json.load(open(ROOT / f"data/questions/json/{qid}.json", encoding="utf-8"))
        s, o, a = q["source"], q["original"], q["answer"]
        src_png = ROOT / o["question_image"]
        src_hash[qid] = sha256(src_png)
        shutil.copyfile(src_png, OUT / "images" / f"{qid}.png")
        rows.append({
            "question_id": qid,
            "subject": s["subject"],
            "exam_year": s["exam_year"],
            "academic_year": s["academic_year"],
            "exam_session": s["exam_type"],
            "exam_name": EXAM_PERIOD.get(s["exam_type"]),
            "exam_date": None,
            "question_number": s["question_number"],
            "stem": o["stem"],
            "passage": o["passage"],
            "choices": o["choices"],
            "question_text": o["raw_text"],
            "official_answer": a["value"],
            "points": o["points"],
            "source_pdf": s["split_pdf"],
            "source_page": s["split_pdf_page"],
            "source_compilation": s["original_compilation"],
            "original_page": s["original_compilation_page"],
            "image": f"images/{qid}.png",
            "image_sha256": src_hash[qid],
            "final_asset_source": sel["source"],
            "review_required": sel["review_required"],
            "primary_curriculum_code": sel["primary_curriculum_code"],
            "kice_sample_link": sel["kice_sample_link"],
        })

    with open(OUT / "questions.jsonl", "w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")

    wb = Workbook()
    ws = wb.active
    ws.title = "LIFE_ETHICS_V2"
    head = ["question_id", "exam_year", "exam_session", "question_number", "official_answer", "points",
            "final_asset_source", "review_required", "primary_curriculum_code", "kice_sample_link", "image_filename"]
    ws.append(head)
    for r in rows:
        link = r["kice_sample_link"]
        ws.append([r["question_id"], r["exam_year"], r["exam_session"], r["question_number"], r["official_answer"],
                   r["points"], r["final_asset_source"], r["review_required"], r["primary_curriculum_code"],
                   f"{link['sample_question']} ({link['match_depth']})" if link else None, r["image"]])
    for i, w in enumerate([32, 10, 12, 10, 10, 7, 28, 10, 18, 22, 44], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(OUT / "questions.xlsx")

    jl = [json.loads(l) for l in open(OUT / "questions.jsonl", encoding="utf-8")]
    pngs = sorted((OUT / "images").glob("*.png"))
    manifest = {
        "package": "LIFE_ETHICS_FINAL_TRANSFER_V2",
        "subject": "LIFE_ETHICS",
        "selection_source": "data/phase2/LIFE_ETHICS/final_content_asset_index.jsonl",
        "total_questions": len(jl),
        "composition": dict(sorted(collections.Counter(r["final_asset_source"] for r in jl).items())),
        "review_required_count": sum(r["review_required"] for r in jl),
        "official_answer_null_count": sum(r["official_answer"] is None for r in jl),
        "image_count": len(pngs),
        "selection_history": {
            "phase2c_keep": sum(1 for d in cur.values() if d == "KEEP"),
            "official_curriculum_core": sum(1 for r in jl if r["final_asset_source"] == "CURRICULUM_EXPLICIT_CORE"),
            "kice_2028_restored": sum(1 for d in rc.values() if d == "RESTORE"),
            "final": len(jl),
        },
        "files": {"questions": "questions.jsonl", "spreadsheet": "questions.xlsx", "images": "images/"},
        "notes": "stem/passage/choices/question_text는 Phase 1 OCR(보조). 원문 Source of Truth는 images/<question_id>.png. "
                 "official_answer는 Phase 1 값 그대로(null은 추론 없이 유지). V1(FINAL_TRANSFER/LIFE_ETHICS)과 별도 패키지.",
        "status": "FROZEN",
    }
    json.dump(manifest, open(OUT / "manifest.json", "w", encoding="utf-8"), ensure_ascii=False, indent=2)

    # ---- QA ----
    xl_rows = openpyxl.load_workbook(OUT / "questions.xlsx").active.max_row - 1
    ids_jl = [r["question_id"] for r in jl]
    ids_png = [p.stem for p in pngs]
    db_ids = {r[0] for r in sqlite3.connect(ROOT / "data/database/questions.db").execute("select id from questions")}
    img_targets = [r["image"] for r in jl]
    qa = {
        "A_index_count": len(idx),
        "B_unique_ids": len(set(ids_idx)),
        "C_core": comp["CURRICULUM_EXPLICIT_CORE"],
        "D_restored": comp["KICE_2028_RESTORED"],
        "E_jsonl_records": len(jl),
        "F_exported_png": len(pngs),
        "G_missing_png": sum(not (OUT / p).exists() for p in img_targets),
        "H_duplicate_png_mapping": len(img_targets) - len(set(img_targets)),
        "I_jsonl_ids_eq_index": set(ids_jl) == set(ids_idx) and len(ids_jl) == len(set(ids_jl)),
        "J_png_ids_eq_index": set(ids_png) == set(ids_idx),
        "K_excluded_contamination": len(set(ids_jl) & excluded) + len(set(ids_png) & excluded)
        + sum(1 for i in ids_jl if cur.get(i) != "KEEP" or rc.get(i) == "DROP"),
        "L_png_sha256_mismatch": sum(sha256(OUT / "images" / f"{i}.png") != src_hash[i] for i in ids_idx),
        "M_xlsx_rows": xl_rows,
        "N_manifest_total": manifest["total_questions"],
        "not_in_phase1_db": sum(i not in db_ids for i in ids_jl),
    }
    ok = (qa["A_index_count"] == qa["B_unique_ids"] == qa["E_jsonl_records"] == qa["F_exported_png"]
          == qa["M_xlsx_rows"] == qa["N_manifest_total"] == 136 and qa["C_core"] == 106 and qa["D_restored"] == 30
          and qa["I_jsonl_ids_eq_index"] and qa["J_png_ids_eq_index"]
          and qa["G_missing_png"] == qa["H_duplicate_png_mapping"] == qa["K_excluded_contamination"]
          == qa["L_png_sha256_mismatch"] == qa["not_in_phase1_db"] == 0
          and manifest["composition"] == {"CURRICULUM_EXPLICIT_CORE": 106, "KICE_2028_RESTORED": 30})
    if not ok:
        raise SystemExit(f"STOP: QA failed {qa}")

    if ZIP.exists():
        ZIP.unlink()
    with zipfile.ZipFile(ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        for p in sorted(OUT.rglob("*")):
            if p.is_file():
                z.write(p, p.relative_to(BASE))

    # ---- ZIP 재오픈 독립 검증 ----
    with zipfile.ZipFile(ZIP) as z:
        bad = z.testzip()
        names = z.namelist()
        roots = {n.split("/")[0] for n in names}
        zjl = [json.loads(l) for l in z.read("LIFE_ETHICS_V2/questions.jsonl").decode("utf-8").splitlines() if l.strip()]
        zimgs = [n for n in names if n.startswith("LIFE_ETHICS_V2/images/") and n.endswith(".png")]
        zip_png_mismatch = sum(hashlib.sha256(z.read(n)).hexdigest() != src_hash[Path(n).stem] for n in zimgs)
        zqa = {
            "readable": True,
            "corrupt_files": 0 if bad is None else 1,
            "root_dirs": sorted(roots),
            "has_questions_jsonl": "LIFE_ETHICS_V2/questions.jsonl" in names,
            "has_questions_xlsx": "LIFE_ETHICS_V2/questions.xlsx" in names,
            "has_manifest": "LIFE_ETHICS_V2/manifest.json" in names,
            "images_count": len(zimgs),
            "jsonl_records": len(zjl),
            "png_sha256_mismatch": zip_png_mismatch,
            "entries": len(names),
        }
    if not (zqa["corrupt_files"] == 0 and zqa["root_dirs"] == ["LIFE_ETHICS_V2"] and zqa["has_questions_jsonl"]
            and zqa["has_questions_xlsx"] and zqa["has_manifest"] and zqa["images_count"] == 136
            and zqa["jsonl_records"] == 136 and zqa["png_sha256_mismatch"] == 0):
        raise SystemExit(f"STOP: ZIP verification failed {zqa}")

    print(json.dumps({"manifest": {k: manifest[k] for k in manifest if k not in ("notes", "files")}, "qa": qa,
                      "zip_qa": zqa, "zip": str(ZIP.relative_to(ROOT)), "zip_bytes": ZIP.stat().st_size},
                     ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
