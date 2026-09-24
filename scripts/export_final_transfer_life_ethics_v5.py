"""LIFE_ETHICS 최종 이관 패키지 V5 export (분석·재판정 없음, V1~V4 보존, 결정적 출력).

선정(유일한 source): data/phase2/LIFE_ETHICS/v4_strict_audit/v4_final_content_asset_index_patched.jsonl (317)
원문: Phase 1 data/questions/json/*.json, data/questions/images/*.png (읽기 전용, PNG byte 그대로 복사)
provenance: v4_strict_audit.jsonl(V4 KEEP 114), v4_targeted_recovery.jsonl(RESTORE 203) — 값 그대로 전달
출력: FINAL_TRANSFER/LIFE_ETHICS_V5/{questions.jsonl, questions.xlsx, images/, manifest.json, README.md}
      FINAL_TRANSFER/LIFE_ETHICS_FINAL_TRANSFER_V5.zip (루트 LIFE_ETHICS_V5/)
결정성: 타임스탬프를 쓰지 않고 xlsx 문서 속성·zip 엔트리 시각을 고정해 재실행 시 byte-identical.
"""
from __future__ import annotations

import collections
import datetime as dt
import hashlib
import io
import json
import re
import shutil
import zipfile
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
P2 = ROOT / "data/phase2/LIFE_ETHICS"
AUD = P2 / "v4_strict_audit"
INDEX = AUD / "v4_final_content_asset_index_patched.jsonl"
MASTER = ROOT / "reference/curriculum/integrated_social_official_scope.json"
BASE = ROOT / "FINAL_TRANSFER"
OUT = BASE / "LIFE_ETHICS_V5"
ZIP = BASE / "LIFE_ETHICS_FINAL_TRANSFER_V5.zip"
EXPECTED = 317
SOURCE_COMMIT = "348744d"  # targeted recovery commit (selection source가 커밋된 시점)
EXAM_PERIOD = {"JUNE": "6월 모의평가", "SEPTEMBER": "9월 모의평가", "CSAT": "대학수학능력시험"}
AREA_KO = {"HAPPINESS": "행복", "ENVIRONMENT": "인간과 자연(환경)", "CIVIL_DISOBEDIENCE": "시민 불복종", "JUSTICE": "사회 정의",
           "PUNISHMENT": "교정적 정의(형벌)", "PEACE": "평화", "AID": "해외 원조·세계 시민", "CULTURE": "문화·다문화"}
FIXED_DT = (2026, 1, 1, 0, 0, 0)


def sha256(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def jl(p):
    return [json.loads(l) for l in open(p, encoding="utf-8")]


def normalized_zip(entries):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries:
            zi = zipfile.ZipInfo(name, date_time=FIXED_DT)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            z.writestr(zi, data)
    return buf.getvalue()


def choices_text(ch):
    return "\n".join(f"{c.get('number')}. {c.get('text')}" for c in (ch or []))


def main():
    idx = jl(INDEX)
    ids = [r["question_id"] for r in idx]
    if len(idx) != EXPECTED or len(set(ids)) != EXPECTED:
        raise SystemExit(f"STOP: selection count {len(idx)} unique {len(set(ids))} != {EXPECTED}")
    sel = {r["question_id"]: r for r in idx}
    v4a = {r["question_id"]: r for r in jl(AUD / "v4_strict_audit.jsonl")}
    rec = {r["question_id"]: r for r in jl(AUD / "v4_targeted_recovery.jsonl")}
    remain_drop = {q for q, r in rec.items() if r["decision"] == "REMAIN_DROP"}
    area_of = {s["code"]: a["area_name"] for a in json.load(open(MASTER, encoding="utf-8"))["areas"] for s in a["standards"]}
    all_ids = {p.stem for p in (ROOT / "data/questions/json").glob("*.json")}

    if OUT.exists():
        shutil.rmtree(OUT)
    (OUT / "images").mkdir(parents=True)

    rows, src_hash, phase1 = [], {}, {}
    for qid in sorted(ids):
        s_ = sel[qid]
        q = json.load(open(ROOT / f"data/questions/json/{qid}.json", encoding="utf-8"))
        phase1[qid] = q
        s, o, a = q["source"], q["original"], q["answer"]
        data = (ROOT / o["question_image"]).read_bytes()
        src_hash[qid] = sha256(data)
        (OUT / "images" / f"{qid}.png").write_bytes(data)
        restored = s_.get("recovery_status") == "RESTORED"
        code = s_.get("primary_curriculum_code")
        if restored:
            r = rec[qid]
            assert r["decision"] == "RESTORE", qid
            prov = {"asset_origin": "V4_TARGETED_RESTORE", "reuse_scope": r["reuse_scope"],
                    "integrated_social_domain": AREA_KO[r["target_area"]], "integrated_social_concept": r["core_content"],
                    "usable_content": r["통합사회에서_활용가능한_내용"], "rationale": r["rationale"],
                    "evidence": {"profile": r["evidence_profile"], "source_evidence": r["source_evidence"]},
                    "philosophers": r["philosophers"], "previous_decision": r["previous_decision"],
                    "review_reasons": r["review_reasons"], "adaptation_required": None}
        else:
            r = v4a[qid]
            assert r["final_decision"] == "KEEP", qid
            prov = {"asset_origin": "V4_KEEP", "reuse_scope": None,
                    "integrated_social_domain": area_of.get(code), "integrated_social_concept": r["core_answer_logic"],
                    "usable_content": None, "rationale": r["rationale"],
                    "evidence": {"official_evidence": r["official_evidence"]},
                    "philosophers": [], "previous_decision": r["previous_status"],
                    "review_reasons": r["review_reasons"], "adaptation_required": r["adaptation_required"]}
        rows.append({
            "question_id": qid,
            "subject": s["subject"],
            "year": s["academic_year"],
            "exam_year": s["exam_year"],
            "exam_type": s["exam_type"],
            "exam_name": EXAM_PERIOD.get(s["exam_type"]),
            "exam_key": s["exam_key"],
            "question_number": s["question_number"],
            "stem": o["stem"],
            "passage": o["passage"],
            "choices": o["choices"],
            "question_text": o["raw_text"],
            "points": o["points"],
            "official_answer": a["value"],
            "source_pdf": s["split_pdf"],
            "source_page": s["split_pdf_page"],
            "source_compilation": s["original_compilation"],
            "original_page": s["original_compilation_page"],
            "image": f"images/{qid}.png",
            "image_sha256": src_hash[qid],
            "review_required": bool(s_["review_required"]),
            "review_reasons": prov.pop("review_reasons"),
            "asset_origin": prov.pop("asset_origin"),
            "reuse_scope": prov.pop("reuse_scope"),
            "integrated_social_domain": prov.pop("integrated_social_domain"),
            "integrated_social_concept": prov.pop("integrated_social_concept"),
            "primary_curriculum_code": code,
            "kice_sample_link": s_.get("kice_sample_link"),
            **prov,
        })

    with open(OUT / "questions.jsonl", "w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")

    wb = Workbook()
    wb.properties.creator = "LIFE_ETHICS_V5"
    wb.properties.created = wb.properties.modified = dt.datetime(*FIXED_DT)
    ws = wb.active
    ws.title = "LIFE_ETHICS_V5"
    head = ["question_id", "year", "exam", "question_number", "stem", "passage", "choices", "points", "official_answer",
            "review_required", "asset_origin", "reuse_scope", "integrated_social_domain", "integrated_social_concept",
            "primary_curriculum_code", "kice_sample_link", "image_filename"]
    ws.append(head)
    for r in rows:
        link = r["kice_sample_link"]
        ws.append([r["question_id"], r["year"], r["exam_name"], r["question_number"], r["stem"], r["passage"],
                   choices_text(r["choices"]), r["points"], r["official_answer"], r["review_required"], r["asset_origin"],
                   r["reuse_scope"], r["integrated_social_domain"], r["integrated_social_concept"], r["primary_curriculum_code"],
                   link.get("sample_question") if link else None, r["image"]])
    for i, w in enumerate([32, 7, 16, 9, 50, 60, 50, 7, 10, 10, 20, 10, 22, 50, 16, 10, 44], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
    for row in ws.iter_rows(min_row=2):
        for i in (4, 5, 6, 13):
            row[i].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    raw = io.BytesIO()
    wb.save(raw)
    fixed_w3c = dt.datetime(*FIXED_DT).strftime("%Y-%m-%dT%H:%M:%SZ")
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as zx:
        parts = []
        for i in zx.infolist():
            data = zx.read(i.filename)
            if i.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                              lambda m: m.group(1) + fixed_w3c.encode() + m.group(2), data)
            parts.append((i.filename, data))
        (OUT / "questions.xlsx").write_bytes(normalized_zip(parts))

    jlr = jl(OUT / "questions.jsonl")
    pngs = sorted((OUT / "images").glob("*.png"))
    origin = collections.Counter(r["asset_origin"] for r in jlr)
    scope = collections.Counter(r["reuse_scope"] for r in jlr if r["reuse_scope"])
    manifest = {
        "package": "LIFE_ETHICS_FINAL_TRANSFER_V5",
        "subject": "LIFE_ETHICS",
        "version": "V5",
        "selection_source": str(INDEX.relative_to(ROOT)),
        "total_questions": len(jlr),
        "original_v4_keep": origin["V4_KEEP"],
        "targeted_restore": origin["V4_TARGETED_RESTORE"],
        "reuse_scope": {"FULL": scope["FULL"], "PARTIAL": scope["PARTIAL"]},
        "review_required_count": sum(r["review_required"] for r in jlr),
        "official_answer_null_count": sum(r["official_answer"] is None for r in jlr),
        "image_count": len(pngs),
        "domain_counts": dict(sorted(collections.Counter(r["integrated_social_domain"] or "기타" for r in jlr).items())),
        "export_timestamp_policy": "타임스탬프 미기록. xlsx 문서 속성과 zip 엔트리 시각을 2026-01-01T00:00:00으로 고정(결정적 출력).",
        "source_commit": SOURCE_COMMIT,
        "recovery_commit": "348744d",
        "exporter": {"path": "scripts/export_final_transfer_life_ethics_v5.py",
                     "sha256": sha256(Path(__file__).read_bytes())},
        "files": {"questions": "questions.jsonl", "spreadsheet": "questions.xlsx", "images": "images/", "readme": "README.md"},
        "previous_packages": "V1(LIFE_ETHICS)·V2·V3·V4는 historical snapshot으로 보존",
        "status": "FROZEN_V5",
    }
    (OUT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "README.md").write_text(
        "# LIFE_ETHICS FINAL TRANSFER V5\n\n"
        "2028 통합사회 모의고사 제작을 위한 생활과 윤리 legacy content source database.\n\n"
        f"- 문항 수: {len(jlr)} = 기존 V4 KEEP {origin['V4_KEEP']} + V4 targeted recovery RESTORE {origin['V4_TARGETED_RESTORE']}\n"
        "- RESTORE는 제공된 통합사회 자료 4개([별책7] 교육과정, 2028 수능 예시문항 안내, 28예시 실제 문항, 최소 성취수준 자료)를 "
        "문항마다 대조한 표적 false-negative recovery 결과이다.\n"
        f"- `reuse_scope`(RESTORE만): FULL {scope['FULL']} / PARTIAL {scope['PARTIAL']} — 원천 콘텐츠 재사용 범위. "
        "PARTIAL은 불량 문항이라는 뜻이 아니라 일부 생윤 고유 세부가 포함되어 있다는 뜻이다.\n"
        "- `asset_origin`: V4_KEEP | V4_TARGETED_RESTORE. `evidence`·`rationale`에 판정 근거를 그대로 보존.\n"
        "- `official_answer`: Life Ethics에는 공식 정답 source가 없어 모두 null(추론·외부 수집 없음).\n"
        "- `questions.jsonl`의 stem/passage/choices는 Phase 1 OCR 값 그대로이며 원문 Source of Truth는 `images/<question_id>.png`.\n"
        "- `review_required`는 후속 제작 시 사람이 주의해서 볼 문항 표시(제외 사유 아님).\n",
        encoding="utf-8")

    # ---- QA ----
    xl = openpyxl.load_workbook(OUT / "questions.xlsx").active
    xl_ids = [row[0] for row in xl.iter_rows(min_row=2, values_only=True)]
    jl_ids = [r["question_id"] for r in jlr]
    png_ids = [p.stem for p in pngs]
    A = set(ids)
    content_mm = {k: 0 for k in ("stem", "passage", "choices", "points", "question_id")}
    for r in jlr:
        o = phase1[r["question_id"]]["original"]
        for k in ("stem", "passage", "choices", "points"):
            content_mm[k] += r[k] != o[k]
        content_mm["question_id"] += r["question_id"] != phase1[r["question_id"]]["question_id"] if "question_id" in phase1[r["question_id"]] else 0
    qa = {
        "selection": len(A), "jsonl": len(jl_ids), "excel_rows": len(xl_ids), "png": len(png_ids),
        "duplicates": (len(jl_ids) - len(set(jl_ids))) + (len(xl_ids) - len(set(xl_ids))) + (len(png_ids) - len(set(png_ids))),
        "missing": len(A - set(jl_ids)) + len(A - set(xl_ids)) + len(A - set(png_ids)),
        "extra": len(set(jl_ids) - A) + len(set(xl_ids) - A) + len(set(png_ids) - A),
        "unknown": len((set(jl_ids) | set(xl_ids) | set(png_ids)) - all_ids),
        "remain_drop_contamination": len((set(jl_ids) | set(png_ids)) & remain_drop),
        "official_answer_non_null": sum(r["official_answer"] is not None for r in jlr),
        "content_mismatch": content_mm,
        "source_to_export_checked": len(png_ids),
        "source_to_export_mismatch": sum(sha256((OUT / "images" / f"{i}.png").read_bytes()) != src_hash[i] for i in ids),
        "restore_without_scope": sum(1 for r in jlr if r["asset_origin"] == "V4_TARGETED_RESTORE" and r["reuse_scope"] not in ("FULL", "PARTIAL")),
    }
    ok = (qa["selection"] == qa["jsonl"] == qa["excel_rows"] == qa["png"] == EXPECTED and set(jl_ids) == set(xl_ids) == set(png_ids) == A
          and all(qa[k] == 0 for k in ("duplicates", "missing", "extra", "unknown", "remain_drop_contamination",
                                       "official_answer_non_null", "source_to_export_mismatch", "restore_without_scope"))
          and not any(content_mm.values()) and origin["V4_KEEP"] == 114 and origin["V4_TARGETED_RESTORE"] == 203)
    if not ok:
        raise SystemExit(f"STOP: QA failed {qa} {dict(origin)}")

    entries = [(str(p.relative_to(BASE)), p.read_bytes()) for p in sorted(OUT.rglob("*")) if p.is_file()]
    ZIP.write_bytes(normalized_zip(entries))
    with zipfile.ZipFile(ZIP) as z:
        bad = z.testzip()
        names = z.namelist()
        zimgs = [n for n in names if n.startswith("LIFE_ETHICS_V5/images/") and n.endswith(".png")]
        zx = openpyxl.load_workbook(io.BytesIO(z.read("LIFE_ETHICS_V5/questions.xlsx"))).active
        zqa = {
            "top_level_dirs": sorted({n.split("/")[0] for n in names}),
            "jsonl_records": sum(1 for l in z.read("LIFE_ETHICS_V5/questions.jsonl").decode("utf-8").splitlines() if l.strip()),
            "excel_rows": zx.max_row - 1,
            "images": len(zimgs),
            "has_manifest_readme": "LIFE_ETHICS_V5/manifest.json" in names and "LIFE_ETHICS_V5/README.md" in names,
            "corrupt_files": 0 if bad is None else 1,
            "source_to_zip_checked": len(zimgs),
            "source_to_zip_mismatch": sum(sha256(z.read(n)) != src_hash[Path(n).stem] for n in zimgs),
        }
    if not (zqa["top_level_dirs"] == ["LIFE_ETHICS_V5"] and zqa["jsonl_records"] == zqa["excel_rows"] == zqa["images"] == EXPECTED
            and zqa["has_manifest_readme"] and zqa["corrupt_files"] == 0 and zqa["source_to_zip_mismatch"] == 0):
        raise SystemExit(f"STOP: ZIP QA failed {zqa}")

    print(json.dumps({"manifest": {k: manifest[k] for k in ("total_questions", "original_v4_keep", "targeted_restore", "reuse_scope",
                                                             "review_required_count", "official_answer_null_count", "image_count")},
                      "qa": qa, "zip_qa": zqa, "zip_bytes": ZIP.stat().st_size, "zip_sha256": sha256(ZIP.read_bytes())},
                     ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
